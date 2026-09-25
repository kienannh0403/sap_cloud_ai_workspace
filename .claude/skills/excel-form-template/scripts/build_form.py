"""build_form.py - FIS "Tools Export Excel" template builder + linter.

Build (spec JSON -> template .xlsx; optional field dictionary .xlsx / ABAP TYPES .abap):
    python build_form.py build <spec.json> --out-dir <dir> [--force] [--with-fields] [--with-types]

Lint an existing template (e.g. one a user already uploaded):
    python build_form.py lint <template.xlsx> [--report ZSD01] [--template ZSD01_EXC]

Spec JSON (UTF-8):
{
  "report":   "ZSD01",                 # Report Name in the upload app (<= 10 chars)
  "template": "ZSD01_EXC",             # Template in the upload app (<= 20 chars)
  "widths":   [20, 40, 10],            # column widths from column B
  "rows": [                            # one entry per template row from row 2; [marker, cellB, cellC, ...]
    ["HEADER", "<NAME>/<BN10LC>-<8>BR"],
    ["",       ""],                    # empty marker = empty output row
    ["ITEM",   "<HDR_CODE>/<BN10CC>Y", "<HDR_NAME>/<BN10CC>Y"],
    ["ITEM_TAB1", "<CODE>/<NN10LC>", "<NAME>/<NN10LC>"]
  ],
  "fields": {                          # optional per-field info (key = field name)
    "NAME":     {"kind": "data",  "source": "I_CompanyCode.CompanyCodeName", "abap_type": "string"},
    "HDR_CODE": {"kind": "label", "text": "Mã"}
  }
}
Rules implemented here mirror references/grammar.md (parser ZBP_R_UPLOAD_TMPL, renderer
ZCL_EXPORT_EXCEL_XLSX, style ZCL_GET_STYLE_TMPL_EXPORT_N).
"""
import argparse
import json
import re
import sys
from collections import OrderedDict
from pathlib import Path

try:
    import openpyxl
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
except ImportError:  # pragma: no cover
    sys.exit("openpyxl is required: pip install openpyxl")

sys.stdout.reconfigure(encoding="utf-8")

MAX_COLS = 100          # ZST_EXCELMAPPING C0001..C0100 (incl. column A)
MAX_ROWS = 255          # ROW_NUM INT1
MAX_MERGE = 255         # MERGE_NGANG / MERGE_DOC INT1
MAX_REPORT = 10         # ZTB_TMPL_EXPORT-REPORT_NAME CHAR 10
MAX_TEMPLATE = 20       # ZTB_UPLOAD_TMPL-TEMPLATE CHAR 20
MAX_LEVEL_VALUE = 10    # ZLEVELVALUE CHAR 10
MAX_FIELD = 30          # ABAP component name

WEIGHTS = {"BN", "BI", "NN", "NI"}
ALIGNS = {"LC", "LT", "LB", "CC", "RC", "RT", "RB"}
COLORS = set("BRYGPWOK")
LABEL_PREFIXES = ("LBL_", "HDR_", "HEADER", "SEC_", "SIGN", "TITLE", "CAP_")

CELL_RE = re.compile(
    r"^<(?P<field>[^<>/]+)>/<(?P<style>[^<>]+)>"
    r"(?P<merges>(?:[-|]<\d+>)*)"
    r"(?P<colors>[A-Z]{0,2})$"
)
MERGE_RE = re.compile(r"([-|])<(\d+)>")
FIELD_OK_RE = re.compile(r"^[A-Z][A-Z0-9_]*$")


class Lint:
    def __init__(self):
        self.errors, self.warnings = [], []

    def err(self, where, msg):
        self.errors.append(f"ERROR   {where}: {msg}")

    def warn(self, where, msg):
        self.warnings.append(f"WARNING {where}: {msg}")

    def report(self):
        for line in self.errors + self.warnings:
            print(line)
        print(f"--- lint: {len(self.errors)} error(s), {len(self.warnings)} warning(s)")
        return not self.errors


def cell_ref(row_idx, col_idx):
    """row_idx: 0-based index in rows (template row = +2); col_idx: 1-based column (A=1)."""
    return f"{get_column_letter(col_idx)}{row_idx + 2}"


def parse_marker(raw, where, lint):
    m = (raw or "").strip().upper()
    if not m:
        return None
    parts = m.split("_")
    layout = parts[0]
    if layout in ("HEADER", "FOOTER") and len(parts) == 1:
        return {"layout": layout, "table": "", "level_field": "", "level_value": ""}
    if layout == "ITEM":
        if len(parts) == 1:
            return {"layout": "ITEM", "table": "", "level_field": "", "level_value": ""}
        if len(parts) == 2 and parts[1]:
            return {"layout": "ITEM", "table": parts[1], "level_field": "", "level_value": ""}
        if len(parts) == 4 and all(parts[1:]):
            if len(parts[3]) > MAX_LEVEL_VALUE:
                lint.err(where, f"level value '{parts[3]}' longer than {MAX_LEVEL_VALUE} chars")
            return {"layout": "ITEM", "table": parts[1], "level_field": parts[2], "level_value": parts[3]}
        lint.err(where, f"marker '{m}' - expected ITEM_<Table> or ITEM_<Table>_<LevelField>_<LevelValue>; "
                        "table/level-field names must not contain '_'")
        return None
    lint.err(where, f"unknown marker '{m}' (HEADER / ITEM / ITEM_<Table>[_<Field>_<Value>] / FOOTER) - never rendered")
    return None


def parse_cell(raw, where, lint):
    v = str(raw).strip().upper()
    m = CELL_RE.match(v)
    if not m:
        lint.err(where, f"'{raw}' does not match <FIELD>/<STYLE>[-<n>][|<n>][COLORS]")
        return None
    field, style, colors = m["field"], m["style"], m["colors"]
    if not FIELD_OK_RE.match(field):
        lint.err(where, f"field '{field}' is not a valid ABAP component name")
    elif len(field) > MAX_FIELD:
        lint.err(where, f"field '{field}' longer than {MAX_FIELD} chars")
    sm = re.fullmatch(r"([A-Z]{2})(\d+)([A-Z]{2})", style)
    if not sm:
        lint.err(where, f"style '{style}' must be <WEIGHT 2 letters><SIZE 2 digits><ALIGN 2 letters>, e.g. BN10LC")
    else:
        weight, size, align = sm.groups()
        if weight not in WEIGHTS:
            lint.err(where, f"weight '{weight}' not in {sorted(WEIGHTS)}")
        if len(size) != 2:
            lint.err(where, f"font size '{size}' must be exactly 2 digits (e.g. 08, 10) - the style class ignores it otherwise")
        if align not in ALIGNS:
            lint.err(where, f"align '{align}' not in {sorted(ALIGNS)} - falls back to Excel default")
    for ch in colors:
        if ch not in COLORS:
            lint.err(where, f"color '{ch}' not in {''.join(sorted(COLORS))}")
    h = v_ = 0
    for kind, n in MERGE_RE.findall(m["merges"]):
        n = int(n)
        if n > MAX_MERGE:
            lint.err(where, f"merge {kind}<{n}> exceeds {MAX_MERGE}")
        if kind == "-":
            h = n
        else:
            v_ = n
    return {"field": field, "style": style, "colors": colors, "merge_h": h, "merge_v": v_}


def analyse(report, template, widths, rows):
    """Lint the grid and collect the field model. rows = [[marker, cellB, ...], ...] starting at template row 2."""
    lint = Lint()
    if report is not None:
        if not report:
            lint.err("meta", "report name is empty")
        elif len(report) > MAX_REPORT:
            lint.err("meta", f"report name '{report}' longer than {MAX_REPORT} chars - upload succeeds but export is empty")
        elif report.upper() == "GEN":
            lint.warn("meta", "report name GEN is the iv_generate base file - it is NOT parsed as a mapping")
    if template is not None:
        if not template:
            lint.err("meta", "template name is empty")
        elif len(template) > MAX_TEMPLATE:
            lint.err("meta", f"template name '{template}' longer than {MAX_TEMPLATE} chars")
    if len(rows) > MAX_ROWS:
        lint.err("grid", f"{len(rows)} template rows (> {MAX_ROWS})")
    width = max([len(r) for r in rows] + [len(widths) + 1, 1])
    if width > MAX_COLS:
        lint.err("grid", f"{width} columns (> {MAX_COLS}: A..CV)")
    for i, w in enumerate(widths):
        if w in ("", None):
            continue
        try:
            float(w)
        except (TypeError, ValueError):
            lint.err(f"{get_column_letter(i + 2)}1", f"width '{w}' is not a number")

    model = {"flat": OrderedDict(), "tables": OrderedDict()}
    occupied = {}                        # (row, col) -> origin ref, for merge-overlap checks
    table_rows = {}                      # table -> list of (level_field, level_value)
    seen_layouts = set()
    for ri, row in enumerate(rows):
        where_row = f"row {ri + 2}"
        marker = parse_marker(row[0] if row else "", f"A{ri + 2}", lint)
        cells = [(ci, c) for ci, c in enumerate(row[1:], start=2) if str(c or "").strip()]
        if marker is None:
            if cells and (row[0] or "").strip() == "":
                lint.warn(where_row, "cells present but column A is empty - the whole row is skipped")
            continue
        seen_layouts.add(marker["layout"])
        if marker["table"]:
            table_rows.setdefault(marker["table"], []).append((marker["level_field"], marker["level_value"]))
        for ci, raw in cells:
            ref = cell_ref(ri, ci)
            c = parse_cell(raw, ref, lint)
            if not c:
                continue
            for dr in range(c["merge_v"] + 1):
                for dc in range(c["merge_h"] + 1):
                    key = (ri + dr, ci + dc)
                    if key in occupied:
                        lint.err(ref, f"{cell_ref(*key)} is covered by both {occupied[key]} and {ref} (merge overlap)")
                    occupied[key] = ref
            if marker["table"]:
                tbl = model["tables"].setdefault(marker["table"], {"level_field": "", "fields": OrderedDict()})
                tbl["fields"].setdefault(c["field"], []).append(ref)
            else:
                area = "ITEM title" if marker["layout"] == "ITEM" else marker["layout"]
                entry = model["flat"].setdefault(c["field"], {"areas": [], "cells": []})
                if area not in entry["areas"]:
                    entry["areas"].append(area)
                entry["cells"].append(ref)

    for tbl, entries in table_rows.items():
        fields = {f for f, _ in entries}
        if len(fields) > 1:
            lint.err(f"ITEM_{tbl}", "mixes leveled and non-leveled rows or several level fields - renderer uses only the first")
        lf = next(iter(fields))
        if tbl in model["tables"]:
            model["tables"][tbl]["level_field"] = lf
            model["tables"][tbl]["level_values"] = [v for _, v in entries if v]
        values = [v for _, v in entries]
        dupes = {v for v in values if values.count(v) > 1}
        if dupes:
            what = "rows" if lf == "" else f"rows for level value(s) {sorted(dupes)}"
            lint.err(f"ITEM_{tbl}", f"duplicate data {what} - only the first is printed, the rest still shift the footer")
        if tbl in model["flat"]:
            lint.err(f"ITEM_{tbl}", f"'{tbl}' is both a table and a flat field")
    if not model["tables"] and "ITEM" in seen_layouts:
        lint.warn("grid", "ITEM title rows but no ITEM_<Table> data row")
    return lint, model


def is_label(name, info):
    kind = (info or {}).get("kind")
    if kind:
        return kind.lower() == "label"
    return name.startswith(LABEL_PREFIXES)


# ---------------------------------------------------------------- outputs
def write_template(path, widths, rows):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "TEMPLATE"
    for i, w in enumerate(widths):
        if w not in ("", None):
            ws.cell(row=1, column=i + 2, value=float(w))
            ws.column_dimensions[get_column_letter(i + 2)].width = float(w)
    for ri, row in enumerate(rows, start=2):
        for ci, val in enumerate(row, start=1):
            if str(val or "").strip():
                c = ws.cell(row=ri, column=ci, value=str(val).strip())
                c.number_format = "@"
                c.font = Font(name="Arial", size=9, bold=(ci == 1))
    ws.column_dimensions["A"].width = 26
    marker_fill = PatternFill("solid", fgColor="9BC2E6")
    for r in range(1, len(rows) + 2):
        ws.cell(row=r, column=1).fill = marker_fill
    ws.freeze_panes = "B2"
    wb.save(path)


def write_fields(path, model, fields_info):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "FIELDS"
    head = ["No", "Field", "Kind", "Area / table", "Template cell(s)", "Label text / data source", "ABAP type"]
    ws.append(head)
    n = 0
    for name, e in model["flat"].items():
        info = fields_info.get(name, {})
        n += 1
        kind = "Label" if is_label(name, info) else "Data"
        ws.append([n, name, kind, ", ".join(e["areas"]), ", ".join(e["cells"]),
                   info.get("text") or info.get("source") or "", info.get("abap_type", "string")])
    for tbl, t in model["tables"].items():
        n += 1
        ws.append([n, tbl, "Table", f"ITEM_{tbl}", "", "nested internal table (1 line = 1 printed row)", f"ty_t_{tbl.lower()}"])
        if t["level_field"]:
            n += 1
            ws.append([n, t["level_field"], "Level", tbl, "",
                       f"selects row format; values {', '.join(t.get('level_values', []))} (UPPERCASE)", "string"])
        for name, cells in t["fields"].items():
            info = fields_info.get(name, {})
            n += 1
            ws.append([n, name, "Data", tbl, ", ".join(cells),
                       info.get("text") or info.get("source") or "", info.get("abap_type", "string")])
    bold = Font(bold=True)
    thin = Side(style="thin", color="999999")
    for c in ws[1]:
        c.font = bold
        c.fill = PatternFill("solid", fgColor="F2DCDB")
    label_fill = PatternFill("solid", fgColor="EAF1FB")
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row):
        for c in row:
            c.border = Border(left=thin, right=thin, top=thin, bottom=thin)
            c.alignment = Alignment(vertical="top", wrap_text=True)
        if row[2].value == "Label":
            for c in row:
                c.fill = label_fill
    for col, w in zip("ABCDEFG", (5, 26, 8, 16, 26, 50, 14)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A2"
    wb.save(path)


def abap_str(text):
    return "`" + str(text).replace("`", "``") + "`"


def write_abap(path, model, fields_info, report, template):
    L = [f'" Generated by excel-form-template - report {report} / template {template}',
         '" 1 row of ty_t_form = 1 sheet. Component names = template tokens.', "TYPES:"]
    for tbl, t in model["tables"].items():
        lo = tbl.lower()
        L.append(f"  BEGIN OF ty_s_{lo},")
        if t["level_field"]:
            vals = ", ".join(t.get("level_values", []))
            L.append(f"    {t['level_field'].lower():<30} TYPE string, \" level field - values {vals} (UPPERCASE)")
        for name in t["fields"]:
            info = fields_info.get(name, {})
            note = f' " {info["source"]}' if info.get("source") else ""
            L.append(f"    {name.lower():<30} TYPE {info.get('abap_type', 'string')},{note}")
        L.append(f"  END OF ty_s_{lo},")
        L.append(f"  ty_t_{lo} TYPE STANDARD TABLE OF ty_s_{lo} WITH EMPTY KEY,")
    L.append("  BEGIN OF ty_s_form,")
    for area in ("HEADER", "ITEM title", "FOOTER"):
        names = [n for n, e in model["flat"].items() if e["areas"][0] == area]
        if not names:
            continue
        L.append(f'    " --- {area} ---')
        for name in names:
            info = fields_info.get(name, {})
            kind = "label" if is_label(name, info) else "data"
            note = info.get("source") or info.get("text") or ""
            L.append(f"    {name.lower():<30} TYPE {info.get('abap_type', 'string')}, \" {kind}{': ' + note if note else ''}")
    if model["tables"]:
        L.append('    " --- item tables (name = ITEM_<Table>) ---')
        for tbl in model["tables"]:
            L.append(f"    {tbl.lower():<30} TYPE ty_t_{tbl.lower()},")
    L.append("  END OF ty_s_form,")
    L.append("  ty_t_form TYPE STANDARD TABLE OF ty_s_form WITH EMPTY KEY.")
    labels = [(n, fields_info.get(n, {}).get("text")) for n, e in model["flat"].items()
              if is_label(n, fields_info.get(n, {}))]
    L += ["", '" Fixed label texts - fill on every sheet row', "DATA(ls_form) = VALUE ty_s_form("]
    L += [f"  {n.lower()} = {abap_str(t) if t else '`TODO`'}" for n, t in labels] or ['  " (no label fields)']
    L[-1] += " )."
    L += ["", "DATA lt_form TYPE ty_t_form.", "APPEND ls_form TO lt_form.", "",
          "DATA lv_file TYPE zde_attachment_tmpl.",
          "NEW zcl_export_excel_xlsx( )->export_excel(",
          f"  EXPORTING iv_template = '{template}'",
          f"            iv_report   = '{report}'",
          "            it_data     = lt_form",
          "            iv_generate = abap_false   \" abap_true needs a GEN base file, else short dump",
          "  IMPORTING lv_context  = lv_file )."]
    Path(path).write_text("\n".join(L) + "\n", encoding="utf-8")


# ---------------------------------------------------------------- commands
def cmd_build(args):
    spec = json.loads(Path(args.spec).read_text(encoding="utf-8"))
    report, template = str(spec.get("report", "")), str(spec.get("template", ""))
    widths, rows = spec.get("widths", []), spec.get("rows", [])
    fields_info = {k.upper(): v for k, v in spec.get("fields", {}).items()}
    lint, model = analyse(report, template, widths, rows)
    unknown = set(fields_info) - set(model["flat"]) - {f for t in model["tables"].values() for f in t["fields"]} \
        - {t["level_field"] for t in model["tables"].values()}
    for name in sorted(unknown):
        lint.warn("fields", f"'{name}' described in spec but not used in any cell")
    ok = lint.report()
    if not ok and not args.force:
        print("NOT written - fix the errors (or --force)")
        return 1
    out = Path(args.out_dir)
    out.mkdir(parents=True, exist_ok=True)
    base = out / template
    write_template(f"{base}.xlsx", widths, rows)
    written = [f"{base}.xlsx"]
    if args.with_fields:
        write_fields(f"{base}_fields.xlsx", model, fields_info)
        written.append(f"{base}_fields.xlsx")
    if args.with_types:
        write_abap(f"{base}_types.abap", model, fields_info, report, template)
        written.append(f"{base}_types.abap")
    print("written: " + ", ".join(written))
    print(f"model: {len(model['flat'])} flat field(s), {len(model['tables'])} table(s): "
          + ", ".join(f"{t}({len(v['fields'])}{', level ' + v['level_field'] if v['level_field'] else ''})"
                      for t, v in model["tables"].items()))
    return 0 if ok else 1


def cmd_lint(args):
    wb = openpyxl.load_workbook(args.xlsx, data_only=True)
    ws = wb.worksheets[0]
    if len(wb.worksheets) > 1:
        print(f"note: {len(wb.worksheets)} sheets - only '{ws.title}' is parsed")
    grid = [[("" if c is None else str(c)) for c in r] for r in ws.iter_rows(values_only=True)]
    widths = grid[0][1:] if grid else []
    rows = grid[1:]
    while rows and not any(x.strip() for x in rows[-1]):
        rows.pop()
    lint, model = analyse(args.report, args.template, widths, rows)
    ok = lint.report()
    print(f"model: {len(model['flat'])} flat field(s), tables: {', '.join(model['tables']) or '-'}")
    return 0 if ok else 1


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = ap.add_subparsers(dest="cmd", required=True)
    b = sub.add_parser("build")
    b.add_argument("spec")
    b.add_argument("--out-dir", required=True)
    b.add_argument("--force", action="store_true")
    b.add_argument("--with-fields", action="store_true", help="also write <template>_fields.xlsx")
    b.add_argument("--with-types", action="store_true", help="also write <template>_types.abap")
    l_ = sub.add_parser("lint")
    l_.add_argument("xlsx")
    l_.add_argument("--report")
    l_.add_argument("--template")
    args = ap.parse_args()
    return cmd_build(args) if args.cmd == "build" else cmd_lint(args)


if __name__ == "__main__":
    sys.exit(main())
